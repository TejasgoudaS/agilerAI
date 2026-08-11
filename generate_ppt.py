import os
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title, bullet_points, notes=""):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title_placeholder.text = title
    
    tf = body_shape.text_frame
    for i, pt in enumerate(bullet_points):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = pt
        p.level = 0
        
    if notes:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes
        
# Slide 1
slide_layout = prs.slide_layouts[0] # Title Slide
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "AI Agile Story Generator"
slide.shapes.placeholders[1].text = "Transforming PRDs into Technical Jira Stories with Multi-Agent AI"
slide.notes_slide.notes_text_frame.text = "Welcome everyone. Today, I'm going to demonstrate the AI Agile Story Generator—a tool we built to solve the tedious and often inaccurate process of manually translating massive Product Requirements Documents into actionable technical tickets in Jira."

# Slide 2
add_slide(
    "The Bottleneck in Agile Planning",
    [
        "Manual Translation is Slow: Product Managers spend hours breaking down PRDs.",
        "Lack of Technical Depth: Standard stories lack architecture or database references.",
        "Traditional AI Falls Short: Single prompts cause hallucinated or generic text."
    ],
    "Currently, translating a PRD into Jira stories is a massive bottleneck. PMs spend hours writing them, and often, engineers find them too vague because they lack technical specifics. We tried using standard AI, but asking one AI to act as an architect, PM, and engineer all at once just causes confusion and hallucinations."
)

# Slide 3
add_slide(
    "Our Solution - The AI 'Crew'",
    [
        "Multi-Agent Orchestration: Powered by CrewAI.",
        "Separation of Concerns: Tasks divided among specialized AI agents.",
        "Sequential Pipeline: Context is passed down an assembly line.",
        "Direct Jira Integration: AI pushes directly to the active sprint."
    ],
    "Our solution is to stop using one AI for everything. Instead, we use a framework called CrewAI to simulate an actual software development team. We have multiple, specialized AI agents working together in an assembly line, and when they are done, the system pushes the results directly into our live Jira workspace."
)

# Slide 4
add_slide(
    "The Technology Stack",
    [
        "The Brain (AI): CrewAI framework + OpenAI's gpt-4o.",
        "The Engine (Backend): Python & FastAPI (real-time SSE streaming).",
        "The Interface (Frontend): React, Vite, and Tailwind CSS.",
        "The Integration: Direct Atlassian Jira Cloud REST API."
    ],
    "Here is a quick look at the tech stack. The heavy AI lifting is done on a Python FastAPI backend using OpenAI's GPT-4o model. The user interacts with a lightning-fast React frontend, which communicates directly with Atlassian Jira's REST API."
)

# Slide 5
add_slide(
    "Meet the AI Agents",
    [
        "1. The Architect Agent: Builds a 'Knowledge Graph' of APIs and databases.",
        "2. The PM Agent: Drafts user stories tied to real tech components.",
        "3. The Engineer Agent: Reviews stories, adjusts complexity, ensures JSON formatting."
    ],
    "Let's meet the team. First, the Architect reads the document and maps out the technical architecture. Second, the PM Agent uses that map to draft stories—ensuring no fake tech is hallucinated. Finally, the Engineer Agent reviews the stories as a quality gate, adjusting story points and adding technical notes."
)

# Slide 6
add_slide(
    "The 'Wow' Factor: Graph RAG",
    [
        "Context over Guesswork: The AI builds a map before writing stories.",
        "Visual Transparency: React Flow renders the AI's 'brain' live.",
        "Zero Hallucinations: Downstream agents must use the Architect's graph."
    ],
    "The secret sauce here is Graph RAG. Before writing anything, the Architect Agent extracts a Knowledge Graph of the system. We actually visualize this live on the screen for the user. Because the PM and Engineer agents are forced to use this graph, the AI cannot hallucinate—it only references the actual services and databases defined in the PRD."
)

# Slide 7
add_slide(
    "Live Demonstration",
    [
        "Connect to live Jira Workspace.",
        "Upload PRD (Client-side parsing).",
        "Watch the real-time AI thought stream.",
        "Review generated stories & Knowledge Graph.",
        "Push to Jira."
    ],
    "Enough talking—let's see it in action. I'm going to connect to our live Jira workspace, upload a PRD, and we will watch the AI crew process the document in real-time."
)

# Slide 8
slide_layout = prs.slide_layouts[0] # Title Slide
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Questions & Answers"
slide.shapes.placeholders[1].text = "Thank You!"
slide.notes_slide.notes_text_frame.text = "Thank you for watching. You've seen how we can take a massive document and turn it into technically accurate Jira tickets in minutes using multi-agent orchestration. I'd love to answer any questions you have about the architecture, CrewAI, or the integration."

output_path = os.path.join(os.getcwd(), "AI_Agile_Story_Generator_Presentation.pptx")
prs.save(output_path)
print(f"Presentation saved to {output_path}")
